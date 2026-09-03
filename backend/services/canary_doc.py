import io
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from backend.core.config import settings
from backend.models.token import HoneyToken


def generate_canary_docx(token: HoneyToken) -> io.BytesIO:
    """
    Generates a Microsoft Word (.docx) document with realistic decoy content
    and an embedded remote web image tracking pixel pointing to the honeytoken beacon.

    When opened in Microsoft Word, LibreOffice, or compatible office suites,
    the application automatically sends a GET request to the beacon URL to fetch
    the remote 1x1 image, logging the accessor's IP and User-Agent.
    """
    doc = Document()

    # Set up realistic decoy document styling
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"CONFIDENTIAL: {token.name.upper()}")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(180, 0, 0)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Classification: Internal Restricted — Level 4 Discretion Required")
    sub_run.font.italic = True
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_heading("1. Executive Summary & Asset Sensitivity", level=2)
    doc.add_paragraph(
        "This document contains privileged operational data, including internal cloud service "
        "credentials, compensation schedules, and proprietary infrastructure topologies. "
        "Unauthorized duplication or transmission constitutes a violation of company policy."
    )

    doc.add_heading("2. Decoy Asset Metadata & Access Tokens", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Parameter"
    hdr_cells[1].text = "Value"

    rows_data = [
        ("Asset Identifier", token.token_id),
        ("Asset Classification", f"Criticality Level {token.criticality} / 50"),
        ("Token Status", token.status.capitalize()),
        ("Placement Target", getattr(token, "placement_location", None) or "Corporate Core Services"),
        ("Access Gate Key", f"AKIA-{token.token_id[:8].upper()}-PROD-SECRET"),
    ]

    for param, val in rows_data:
        row_cells = table.add_row().cells
        row_cells[0].text = param
        row_cells[1].text = val

    doc.add_heading("3. Operational Notes", level=2)
    doc.add_paragraph(
        "All queries against the underlying database clusters must use TLS 1.3 with authorized "
        "mTLS client certificates. For escalations, contact the internal SecOps response team."
    )

    # Construct the beacon tracking URL
    beacon_url = f"{settings.BEACON_BASE_URL.rstrip('/')}/beacon/{token.token_id}"

    # Embed external relationship in document part
    # OpenXML relationship for external image
    rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    rel_id = doc.part.relate_to(beacon_url, rel_type, is_external=True)

    # Construct OpenXML Drawing element with external blip link
    drawing_xml = f"""
    <w:p {nsdecls('w', 'r', 'wp', 'a', 'pic')}>
      <w:pPr>
        <w:pBdr>
          <w:bottom w:val="none"/>
        </w:pBdr>
      </w:pPr>
      <w:r>
        <w:drawing>
          <wp:inline distT="0" distB="0" distL="0" distR="0">
            <wp:extent cx="9525" cy="9525"/>
            <wp:docPr id="101" name="CanaryTrackingPixel"/>
            <a:graphic>
              <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
                <pic:pic>
                  <pic:nvPicPr>
                    <pic:cNvPr id="102" name="PixelImage"/>
                    <pic:cNvPicPr/>
                  </pic:nvPicPr>
                  <pic:blipFill>
                    <a:blip r:link="{rel_id}"/>
                    <a:stretch>
                      <a:fillRect/>
                    </a:stretch>
                  </pic:blipFill>
                  <pic:spPr>
                    <a:xfrm>
                      <a:off x="0" y="0"/>
                      <a:ext cx="9525" cy="9525"/>
                    </a:xfrm>
                    <a:prstGeom prst="rect">
                      <a:avLst/>
                    </a:prstGeom>
                  </pic:spPr>
                </pic:pic>
              </a:graphicData>
            </a:graphic>
          </wp:inline>
        </w:drawing>
      </w:r>
    </w:p>
    """
    p_element = parse_xml(drawing_xml)
    doc._body._element.append(p_element)

    # Save to in-memory byte buffer
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
