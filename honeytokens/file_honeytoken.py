"""
Multi-Format Decoy File Honeytoken Generator
CyberGuard SOC — Honeytoken-Based Insider Threat Detector
"""

import io
import json
import os
import re
import secrets
import socket
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOKEN_FILE = os.path.join(os.path.dirname(__file__), "tokens.json")
DECOY_FOLDER = os.path.join(BASE_DIR, "decoy_files")

os.makedirs(DECOY_FOLDER, exist_ok=True)


def get_lan_ip() -> str:
    """Return this machine's active LAN IP address."""
    sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        sock.connect(("8.8.8.8", 80))
        return sock.getsockname()[0]
    except OSError:
        try:
            return socket.gethostbyname(socket.gethostname())
        except OSError:
            return "127.0.0.1"
    finally:
        sock.close()


def _load_tokens():
    if not os.path.exists(TOKEN_FILE):
        return []
    try:
        with open(TOKEN_FILE, "r", encoding="utf-8") as file:
            return json.load(file)
    except (json.JSONDecodeError, FileNotFoundError):
        return []


def _save_tokens(tokens):
    with open(TOKEN_FILE, "w", encoding="utf-8") as file:
        json.dump(tokens, file, indent=4)


def _make_unique_filename(filename, tokens):
    """Return a safe, unique decoy filename."""
    filename = os.path.basename(filename).strip()
    if not filename:
        filename = "Executive_Payroll_2026.pdf"

    # Remove characters unsafe in Windows filenames
    filename = re.sub(r'[<>:"/\\|?*]', "_", filename)
    filename = filename.rstrip(". ")

    stem, extension = os.path.splitext(filename)
    if not extension:
        extension = ".pdf"

    existing = {str(t.get("filename", "")).lower() for t in tokens}
    candidate = f"{stem}{extension}"
    number = 1

    while candidate.lower() in existing or os.path.exists(
        os.path.join(DECOY_FOLDER, candidate)
    ):
        candidate = f"{stem}_{number:03d}{extension}"
        number += 1

    return candidate


def build_canary_pdf_bytes(token_id: str, filename: str, server_ip: str) -> bytes:
    """
    Generate valid binary PDF-1.4 bytes containing:
    1. /OpenAction URI pointing to http://{server_ip}:5000/beacon/{token_id}
    2. Visible confidential financial / payroll decoy text stream
    """
    beacon_url = f"http://{server_ip}:5000/beacon/{token_id}"
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    stream_content = (
        "BT\n"
        "/F1 18 Tf\n"
        "50 740 Td\n"
        f"(CYBERGUARD SOC - CONFIDENTIAL INTERNAL PAYROLL) Tj\n"
        "/F1 10 Tf\n"
        "0 -25 Td\n"
        f"(Document: {filename} | Generated: {now_str}) Tj\n"
        "0 -25 Td\n"
        f"(Tracking Asset ID: {token_id}) Tj\n"
        "0 -20 Td\n"
        "(Classification: RESTRICTED - EXECUTIVE BOARD ONLY) Tj\n"
        "0 -30 Td\n"
        "/F1 12 Tf\n"
        "(Executive Compensation Matrix FY2026-27:) Tj\n"
        "0 -20 Td\n"
        "( - CEO: Base Salary $450,000 | Stock Options: 120,000 Shares) Tj\n"
        "0 -18 Td\n"
        "( - CTO: Base Salary $380,000 | Stock Options: 95,000 Shares) Tj\n"
        "0 -18 Td\n"
        "( - CFO: Base Salary $360,000 | Stock Options: 90,000 Shares) Tj\n"
        "0 -35 Td\n"
        "/F1 10 Tf\n"
        "(Security Notice: Access to this file is cryptographically monitored and audited.) Tj\n"
        "ET\n"
    ).encode("utf-8")

    stream_len = len(stream_content)
    pdf_parts = []
    offsets = []

    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    pdf_parts.append(header)

    def add_object(obj_bytes: bytes) -> int:
        nonlocal pdf_parts, offsets
        offset = sum(len(p) for p in pdf_parts)
        offsets.append(offset)
        pdf_parts.append(obj_bytes)
        return len(offsets)

    # 1 0 obj: Catalog with OpenAction
    add_object(
        b"1 0 obj\n"
        b"<<\n"
        b"  /Type /Catalog\n"
        b"  /Pages 2 0 R\n"
        b"  /OpenAction 5 0 R\n"
        b">>\n"
        b"endobj\n"
    )

    # 2 0 obj: Pages
    add_object(
        b"2 0 obj\n"
        b"<<\n"
        b"  /Type /Pages\n"
        b"  /Kids [3 0 R]\n"
        b"  /Count 1\n"
        b">>\n"
        b"endobj\n"
    )

    # 3 0 obj: Page
    add_object(
        b"3 0 obj\n"
        b"<<\n"
        b"  /Type /Page\n"
        b"  /Parent 2 0 R\n"
        b"  /MediaBox [0 0 612 792]\n"
        b"  /Contents 4 0 R\n"
        b"  /Resources <<\n"
        b"    /Font <<\n"
        b"      /F1 <<\n"
        b"        /Type /Font\n"
        b"        /Subtype /Type1\n"
        b"        /BaseFont /Helvetica\n"
        b"      >>\n"
        b"    >>\n"
        b"  >>\n"
        b"  /Annots [6 0 R]\n"
        b">>\n"
        b"endobj\n"
    )

    # 4 0 obj: Page contents
    obj4 = (
        f"4 0 obj\n<< /Length {stream_len} >>\nstream\n".encode("utf-8")
        + stream_content
        + b"\nendstream\nendobj\n"
    )
    add_object(obj4)

    # 5 0 obj: OpenAction pointing to Honeytoken beacon URI
    safe_beacon_url = beacon_url.replace("(", "\\(").replace(")", "\\)")
    obj5 = (
        f"5 0 obj\n"
        f"<<\n"
        f"  /Type /Action\n"
        f"  /S /URI\n"
        f"  /URI ({safe_beacon_url})\n"
        f">>\n"
        f"endobj\n"
    ).encode("utf-8")
    add_object(obj5)

    # 6 0 obj: Annot Link
    obj6 = (
        f"6 0 obj\n"
        f"<<\n"
        f"  /Type /Annot\n"
        f"  /Subtype /Link\n"
        f"  /Rect [50 710 500 760]\n"
        f"  /Border [0 0 0]\n"
        f"  /A 5 0 R\n"
        f">>\n"
        f"endobj\n"
    ).encode("utf-8")
    add_object(obj6)

    # Xref Table
    xref_offset = sum(len(p) for p in pdf_parts)
    xref = [f"xref\n0 {len(offsets) + 1}\n0000000000 65535 f \n".encode("ascii")]
    for off in offsets:
        xref.append(f"{off:010d} 00000 n \n".encode("ascii"))
    pdf_parts.extend(xref)

    # Trailer
    trailer = (
        f"trailer\n"
        f"<<\n"
        f"  /Size {len(offsets) + 1}\n"
        f"  /Root 1 0 R\n"
        f">>\n"
        f"startxref\n"
        f"{xref_offset}\n"
        f"%%EOF\n"
    ).encode("ascii")
    pdf_parts.append(trailer)

    return b"".join(pdf_parts)


def build_canary_docx(token_id: str, filename: str, server_ip: str, file_path: str):
    """
    Generate a realistic Word document with an external tracking relationship
    pointing to http://{server_ip}:5000/beacon/{token_id}.
    """
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Pt, RGBColor

    doc = Document()
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"CONFIDENTIAL STRATEGY: {filename}")
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(180, 0, 0)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Classification: Internal Restricted — Board of Directors Discretion Only")
    sub_run.font.italic = True
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_heading("1. Strategic M&A Advisory & Market Expansion", level=2)
    doc.add_paragraph(
        "This internal strategy document outlines projected corporate acquisition targets "
        "and proprietary financial hedges for Q3/Q4. All access events are registered."
    )

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Security Parameter"
    hdr_cells[1].text = "Designation"

    rows = [
        ("Tracking Token ID", token_id),
        ("Document Classification", "Restricted Corporate Asset"),
        ("Active Defense Beacon", f"http://{server_ip}:5000/beacon/{token_id}"),
        ("Audit Verification", "Centralized SOC Logging Active"),
    ]
    for param, val in rows:
        r_cells = table.add_row().cells
        r_cells[0].text = param
        r_cells[1].text = val

    # Embed external relationship for image beacon
    beacon_url = f"http://{server_ip}:5000/beacon/{token_id}"
    rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    rel_id = doc.part.relate_to(beacon_url, rel_type, is_external=True)

    drawing_xml = f"""
    <w:p {nsdecls('w', 'r', 'wp', 'a', 'pic')}>
      <w:r>
        <w:drawing>
          <wp:inline distT="0" distB="0" distL="0" distR="0">
            <wp:extent cx="9525" cy="9525"/>
            <wp:docPr id="101" name="CyberGuardBeaconPixel"/>
            <a:graphic>
              <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
                <pic:pic>
                  <pic:nvPicPr>
                    <pic:cNvPr id="102" name="BeaconPixel"/>
                    <pic:cNvPicPr/>
                  </pic:nvPicPr>
                  <pic:blipFill>
                    <a:blip r:link="{rel_id}"/>
                    <a:stretch><a:fillRect/></a:stretch>
                  </pic:blipFill>
                  <pic:spPr>
                    <a:xfrm><a:off x="0" y="0"/><a:ext cx="9525" cy="9525"/></a:xfrm>
                    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                  </pic:spPr>
                </pic:pic>
              </a:graphicData>
            </a:graphic>
          </wp:inline>
        </w:drawing>
      </w:r>
    </w:p>
    """
    doc._body._body.append(parse_xml(drawing_xml))
    doc.save(file_path)


def build_credential_decoy_text(token_id: str, server_ip: str, ext: str) -> str:
    """Generate realistic cloud credentials with embedded honeytoken verification endpoints."""
    aws_key = "AKIA" + secrets.token_hex(8).upper()
    aws_secret = secrets.token_urlsafe(32)
    stripe_key = "sk_live_" + secrets.token_urlsafe(24)

    if ext in [".yaml", ".yml"]:
        return f"""# Production Cloud Services Configuration
# Classification: STRICTLY CONFIDENTIAL - DO NOT COMMIT TO PUBLIC REPO
environment: production
cloud_provider: aws
region: us-east-1

aws:
  access_key_id: "{aws_key}"
  secret_access_key: "{aws_secret}"

payment_gateway:
  provider: stripe
  live_secret_key: "{stripe_key}"

internal_endpoints:
  auth_verify_uri: "http://{server_ip}:5000/api/v1/auth/verify"
  beacon_health_uri: "http://{server_ip}:5000/beacon/{token_id}"
  telemetry_id: "{token_id}"
"""
    # .env format
    return f"""# Production Environment Secrets & API Keys
# CONFIDENTIAL - INTERNAL USE ONLY
ENVIRONMENT=production
AWS_REGION=us-east-1
AWS_ACCESS_KEY_ID={aws_key}
AWS_SECRET_ACCESS_KEY={aws_secret}

STRIPE_LIVE_SECRET={stripe_key}

# Internal Gateway Verification & Health Checks
INTERNAL_AUTH_VERIFY_ENDPOINT=http://{server_ip}:5000/api/v1/auth/verify
BEACON_HEALTH_URI=http://{server_ip}:5000/beacon/{token_id}
HONEYTOKEN_ID={token_id}
"""


def generate_file_honeytoken(filename="Executive_Payroll_2026.pdf", server_ip=None):
    """
    Generate a decoy honeytoken file across multiple formats (.pdf, .docx, .env, .yaml, .txt).
    Embeds canary traps linking back to the CyberGuard beacon and honeypot gateway.
    """
    tokens = _load_tokens()
    filename = _make_unique_filename(filename, tokens)
    token_id = secrets.token_hex(16)
    sip = server_ip or get_lan_ip()

    token = {
        "token_id": token_id,
        "token_type": "FILE",
        "filename": filename,
        "created_at": datetime.now().isoformat(),
        "status": "ACTIVE",
        "trigger_url": f"/beacon/{token_id}",
    }

    tokens.append(token)
    _save_tokens(tokens)

    file_path = os.path.join(DECOY_FOLDER, filename)
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        pdf_bytes = build_canary_pdf_bytes(token_id, filename, sip)
        with open(file_path, "wb") as f:
            f.write(pdf_bytes)
    elif ext == ".docx":
        try:
            build_canary_docx(token_id, filename, sip, file_path)
        except Exception:
            # Fallback if docx construction errors
            with open(file_path, "wb") as f:
                f.write(build_canary_pdf_bytes(token_id, filename, sip))
    elif ext in [".env", ".yaml", ".yml"]:
        content = build_credential_decoy_text(token_id, sip, ext)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)
    else:
        # Default text / report format
        content = f"""============================================================
CYBERGUARD SOC — CONFIDENTIAL RESOURCE AUDIT
============================================================
Document: {filename}
Token ID: {token_id}
Classification: RESTRICTED ACCESS
Timestamp: {datetime.now().isoformat()}

INTERNAL AUDIT TELEMETRY BEACON:
http://{sip}:5000/beacon/{token_id}

Notice: All accesses to this document are audited and analyzed
in real time by CyberGuard SOC.
============================================================
"""
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)

    return token
